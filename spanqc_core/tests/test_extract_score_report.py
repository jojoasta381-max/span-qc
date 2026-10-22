import json
from conftest import FakeClient, reply, wire
from spanqc_core.extract import ModelReply, extract_sheet, parse_reply
from spanqc_core.score import aggregate, score_sheet, write_run_log
from spanqc_core.wirecodes import WireLabel
from spanqc_core.classify import ComponentIndex
from spanqc_core.report import build_report
from spanqc_core.rules import Observation


def test_happy_path():
    r = extract_sheet(FakeClient([reply([wire()])]), b"x", "image/png")
    assert r.ok and r.observations[0].wire_no == "122A" and r.observations[0].confidence == 0.95


def test_truncation_fails_the_page_and_trusts_nothing():
    r = extract_sheet(FakeClient([reply([wire()], stop="max_tokens")]), b"x", "image/png")
    assert not r.ok and r.truncated and r.observations == []


def test_garbage_json_fails_the_page():
    assert not extract_sheet(FakeClient([ModelReply("sorry I cannot", "end_turn")]), b"x", "image/png").ok


def test_markdown_fenced_json_accepted():
    body = "```json\n" + json.dumps({"wires": [wire()], "uncertain": []}) + "\n```"
    assert extract_sheet(FakeClient([ModelReply(body, "end_turn")]), b"x", "image/png").ok


def test_model_exception_fails_the_page_never_fakes():
    r = extract_sheet(FakeClient([RuntimeError("401")]), b"x", "image/png")
    assert not r.ok and r.observations == []


def test_evidence_that_does_not_support_fields_forces_review():
    r = extract_sheet(FakeClient([reply([wire(ev="122A BK 1.5"), wire("7", ev=""), wire("8", ev="8 WH 2.0", size=1.5)])]), b"x", "image/png")
    assert [o.confidence for o in r.observations] == [0.0, 0.0, 0.0]


def test_evidence_with_equivalent_size_formatting_is_accepted():
    r = extract_sheet(FakeClient([reply([wire("8", "BK", 2, ev="8 BK 2.0")])]), b"x", "image/png")
    assert r.observations[0].confidence == 0.95


def test_missing_confidence_defaults_to_zero():
    w = wire(); w.pop("confidence")
    assert extract_sheet(FakeClient([reply([w])]), b"x", "image/png").observations[0].confidence == 0.0


def test_no_value_is_ever_filled_in():
    w = wire(col=None, ev="122A  1.5")
    o = extract_sheet(FakeClient([reply([w])]), b"x", "image/png").observations[0]
    assert o.colour is None


def test_scoring_flags_fabrication(tmp_path):
    truth = [WireLabel("1", "BK", "0.75"), WireLabel("2", "RD", "1.5")]
    s = score_sheet("S", truth, [{"wire_no": "1", "colour": "BK", "size": "0.75"}, {"wire_no": "9", "colour": "WH", "size": "1"}])
    assert (s.matched, len(s.fabricated), len(s.missed)) == (1, 1, 1)
    agg = aggregate([s])
    assert agg["fabricated_count"] == 1 and agg["recall"] == 0.5
    rec = write_run_log(tmp_path / "log.json", prompt_version="v1", model="m", change_note="baseline", scores=[s])
    assert json.loads((tmp_path / "log.json").read_text())[0]["prompt_version"] == "v1" and rec["aggregate"]["precision"] == 0.5


def test_scoring_unparsable_counts_as_fabricated():
    s = score_sheet("S", [], ["garbage", {"wire_no": None}])
    assert s.fabricated and s.observed == 2


def test_report_builds_with_prompt_version(tmp_path):
    from openpyxl import load_workbook
    p = build_report(tmp_path / "r.xlsx", overall="REVIEW", findings=[], observations=[Observation("1", "BK", "1", "CFB1", "KL30", 0.9, "1 BK 1")],
                     index=ComponentIndex(in_scope={"CFB1"}), prompt_version="extract_v1", model="m", document="d.pdf", sha256="ab", pages=(1, 1))
    wb = load_workbook(p)
    assert wb.sheetnames == ["Summary", "Wires", "Findings", "Flagged gaps", "Components"]
    assert "extract_v1" in [c.value for row in wb["Summary"].iter_rows() for c in row]
    assert wb["Wires"]["G2"].value == "blue"


def test_api_error_reason_is_surfaced_and_workspace_header_sent(monkeypatch):
    import httpx
    from spanqc_core.extract import AnthropicClient
    seen = {}

    def fake_post(url, json=None, headers=None, timeout=None):
        seen["headers"] = headers
        return httpx.Response(400, json={"error": {"message": "workspace needed"}}, request=httpx.Request("POST", url))

    monkeypatch.setattr(httpx, "post", fake_post)
    monkeypatch.setenv("ANTHROPIC_WORKSPACE_ID", "wrkspc_123")
    r = extract_sheet(AnthropicClient(api_key="k"), b"x", "image/png")
    assert not r.ok and "workspace needed" in r.error and seen["headers"]["anthropic-workspace-id"] == "wrkspc_123"
