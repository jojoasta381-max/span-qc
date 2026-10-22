"""Multi-tab Excel report with the prompt version recorded (so every result is reproducible)."""
from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Dict, List, Sequence

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .classify import ComponentIndex, DestClass, classify_destination
from .rules import Observation

_FILL = {DestClass.GREEN: "C6EFCE", DestClass.YELLOW: "FFEB9C", DestClass.GREY: "D9D9D9", DestClass.BLUE: "BDD7EE"}
_OUT = {"PASS": "C6EFCE", "REVIEW": "FFEB9C", "FAIL": "FFC7CE"}


def _header(ws, cols):
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")


def _autosize(ws):
    for i, col in enumerate(ws.columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, max(len(str(c.value or "")) for c in col) + 2))


def build_report(path: str | Path, *, overall: str, findings: Sequence[dict], observations: Sequence[Observation],
                 index: ComponentIndex, prompt_version: str, model: str, document: str, sha256: str,
                 pages: tuple[int, int], uncertain: Sequence[dict] = ()) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    classes: Dict[DestClass, int] = {c: 0 for c in DestClass}
    rows = []
    for o in observations:
        cls = classify_destination(o.destination, index)
        classes[cls] += 1
        rows.append((o, cls))
    for k, v in [("Document", document), ("SHA-256", sha256), ("Generated (UTC)", datetime.now(timezone.utc).isoformat(timespec="seconds")),
                 ("Prompt version", prompt_version), ("Model", model), ("Pages processed", f"{pages[0]}/{pages[1]}"),
                 ("Overall (machine)", overall), ("Findings", len(findings)),
                 ("Green / Yellow / Grey / Blue", " / ".join(str(classes[c]) for c in DestClass)),
                 ("Status", "Machine result only. A human reviewer must sign off.")]:
        ws.append([k, v])
    ws["A1"].font = ws["A4"].font = Font(bold=True)
    for r in ws.iter_rows(min_row=1, max_row=ws.max_row):
        r[0].font = Font(bold=True)
    ws.cell(row=7, column=2).fill = PatternFill("solid", fgColor=_OUT.get(overall, "FFFFFF"))
    _autosize(ws)

    wr = wb.create_sheet("Wires")
    _header(wr, ["Page", "Wire", "Colour", "Size mm2", "Source", "Destination", "Dest class", "Confidence", "Evidence"])
    for o, cls in rows:
        wr.append([o.page, o.wire_no, o.colour, o.size, o.source, o.destination, cls.value, round(o.confidence, 2), o.evidence])
        wr.cell(row=wr.max_row, column=7).fill = PatternFill("solid", fgColor=_FILL[cls])
    _autosize(wr)

    wf = wb.create_sheet("Findings")
    _header(wf, ["Rule", "Outcome", "Title", "Reason", "Expected", "Detected", "Page"])
    for f in findings:
        wf.append([f["rule_id"], f["outcome"], f["title"], f["reason"], f["expected_value"], f["detected_value"], f["page_number"]])
        wf.cell(row=wf.max_row, column=2).fill = PatternFill("solid", fgColor=_OUT.get(f["outcome"], "FFFFFF"))
    _autosize(wf)

    wu = wb.create_sheet("Flagged gaps")
    _header(wu, ["What", "Why"])
    for u in uncertain:
        wu.append([u.get("what"), u.get("why")])
    _autosize(wu)

    wc = wb.create_sheet("Components")
    _header(wc, ["Component", "Scope"])
    for c in sorted(index.known):
        wc.append([c, "in scope" if c in index.in_scope else "known (not re-verified)"])
    _autosize(wc)
    out = Path(path)
    wb.save(out)
    return out
